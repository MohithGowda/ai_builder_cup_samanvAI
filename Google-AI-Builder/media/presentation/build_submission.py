from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = Path(r"C:\Users\z004s5sc\Downloads\Submission Template _ AI Builder Cup.pptx")
OUT_DIR = ROOT / "media" / "presentation"
OUT_FILE = OUT_DIR / "SamanvAI_AI_Builder_Cup_Submission.pptx"
CAPTURE_DIR = ROOT / "media" / "demo-video"
SCREEN_DIR = OUT_DIR / "screens"

INK = "1C3038"
MUTED = "526B73"
CYAN = "087F91"
BRIGHT = "48C6D0"
PALE = "E5F2F5"
PALE_GREEN = "EEF6F2"
LINE = "CBDDE1"
AMBER = "A66B24"
PALE_AMBER = "FAF2E7"
WHITE = "FFFFFF"
BLACK = "111416"
FONT = "Arial"


def rgb(value):
    return RGBColor.from_string(value)


def style_run(run, size, color=INK, bold=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(slide, text, x, y, w, h, size=12, color=INK, bold=False, align=None, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(3)
        if align is not None:
            paragraph.alignment = align
        for run in paragraph.runs:
            style_run(run, size, color, bold)
    return box


def add_rich_lines(slide, rows, x, y, w, h, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    for index, row in enumerate(rows):
        text, size, color, bold, after = row
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.space_after = Pt(after)
        for run in paragraph.runs:
            style_run(run, size, color, bold)
    return box


def add_card(slide, x, y, w, h, title, body, accent=CYAN, body_size=10.4, fill=WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = rgb(fill)
    card.line.color.rgb = rgb(LINE)
    card.line.width = Pt(0.8)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = rgb(accent)
    stripe.line.fill.background()
    add_text(slide, title, x + 0.18, y + 0.12, w - 0.30, 0.38, 12.2, INK, True)
    add_text(slide, body, x + 0.18, y + 0.55, w - 0.32, h - 0.65, body_size, MUTED)
    return card


def add_pill(slide, x, y, w, text, fill=PALE, color=CYAN, size=9.4):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    add_text(slide, text, x + 0.04, y + 0.01, w - 0.08, 0.28, size, color, True, PP_ALIGN.CENTER, 0)


def add_note(slide, text, y=5.12, color=MUTED):
    add_text(slide, text, 0.40, y, 9.16, 0.25, 8.2, color, False)


def set_section_title(slide, title):
    text_shape = next((shape for shape in slide.shapes if shape.has_text_frame and shape.text.strip()), None)
    if text_shape is None:
        text_shape = slide.shapes.add_textbox(Inches(0.29), Inches(0.73), Inches(9.78), Inches(0.61))
    frame = text_shape.text_frame
    frame.clear()
    frame.word_wrap = False
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    for run in paragraph.runs:
        style_run(run, 16, "111111", True)
    return text_shape


def body_slide(presentation, number, title):
    slide = presentation.slides[number - 1]
    set_section_title(slide, title)
    return slide


def add_chevron(slide, x, y, w=0.19, h=0.40):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(BRIGHT)
    shape.line.fill.background()


def add_flow_node(slide, x, title, body, width=1.38, fill=PALE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.08), Inches(width), Inches(1.18))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(LINE)
    shape.line.width = Pt(0.8)
    add_text(slide, title, x + 0.08, 2.23, width - 0.16, 0.38, 10.3, CYAN, True, PP_ALIGN.CENTER)
    add_text(slide, body, x + 0.10, 2.65, width - 0.20, 0.49, 8.6, INK, False, PP_ALIGN.CENTER)


def add_arch_node(slide, x, y, w, title, body, fill=PALE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.78))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(LINE)
    shape.line.width = Pt(0.8)
    add_text(slide, title, x + 0.08, y + 0.09, w - 0.16, 0.27, 9.6, CYAN, True, PP_ALIGN.CENTER, 0.01)
    add_text(slide, body, x + 0.08, y + 0.40, w - 0.16, 0.29, 8.0, INK, False, PP_ALIGN.CENTER, 0.01)


def crop_capture(source, destination):
    image = Image.open(source).convert("RGB")
    width, height = image.size
    target_ratio = 16 / 9
    visible_width = int(width * 0.475)
    visible_height = int(height * 0.475)
    if visible_width / visible_height > target_ratio:
        visible_width = int(visible_height * target_ratio)
    else:
        visible_height = int(visible_width / target_ratio)
    image.crop((0, 0, visible_width, visible_height)).save(destination, quality=94)
    return destination


def add_picture_crop(slide, path, x, y, w, h):
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    picture.crop_to_fit = True if hasattr(picture, "crop_to_fit") else False
    picture.line.color.rgb = rgb(LINE)
    picture.line.width = Pt(0.8)
    return picture


def add_bullet_stack(slide, x, y, w, rows, size=11, gap=0.48, accent=CYAN):
    for index, (title, body) in enumerate(rows):
        top = y + index * gap
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(top + 0.08), Inches(0.09), Inches(0.09))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(accent)
        dot.line.fill.background()
        add_text(slide, title, x + 0.18, top, w - 0.18, 0.22, size, INK, True)
        add_text(slide, body, x + 0.18, top + 0.22, w - 0.22, gap - 0.22, size - 2, MUTED)


def build():
    if not TEMPLATE.is_file():
        raise FileNotFoundError(f"Submission template not found: {TEMPLATE}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SCREEN_DIR.mkdir(parents=True, exist_ok=True)
    screenshots = {}
    for key, name in {
        "overview": "01-governance.jpg",
        "map": "02-live-map.jpg",
        "citizen": "03-citizen-home.jpg",
        "assistant": "04-samanvay-guidance.jpg",
        "form": "06-birth-certificate-form.jpg",
        "architecture": "07-cloud-architecture.jpg",
    }.items():
        source = CAPTURE_DIR / name
        if source.is_file():
            screenshots[key] = crop_capture(source, SCREEN_DIR / f"{key}.jpg")

    presentation = Presentation(str(TEMPLATE))
    if len(presentation.slides) != 16:
        raise ValueError(f"Expected 16 template slides, found {len(presentation.slides)}")
    presentation.core_properties.title = "SamanvAI | AI Builder Cup Submission"
    presentation.core_properties.subject = "Civic service access and governance intelligence pilot"
    presentation.core_properties.author = "SamanvAI Project Team"
    presentation.core_properties.keywords = "civic technology, citizen services, governance, LangGraph, Cloud Run"

    cover = presentation.slides[0].shapes[0]
    frame = cover.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    cover_rows = [
        ("Team Details", 21, INK, True, 8),
        ("Team name: SamanvAI", 13, INK, True, 3),
        ("Team lead: [Add team lead name]", 12, MUTED, False, 3),
        ("Problem: citizens navigate fragmented services while officials lack a unified, jurisdiction-scoped operating view.", 11, MUTED, False, 0),
    ]
    for index, (text, size, color, bold, after) in enumerate(cover_rows):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.space_after = Pt(after)
        for run in paragraph.runs:
            style_run(run, size, color, bold)

    overview = presentation.slides[1]
    instruction_shape = overview.shapes[2]
    instruction_shape.left = Inches(5.20)
    instruction_shape.top = Inches(1.22)
    instruction_shape.width = Inches(4.25)
    instruction_shape.height = Inches(3.85)
    frame = instruction_shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    overview_lines = [
        ("One connected civic service loop", 22, "FFFFFF", True, 15),
        ("CITIZENS", 9, "48C6D0", True, 2),
        ("Discover a service, create a demo request, and track its next step.", 13, "FFFFFF", False, 9),
        ("OFFICIALS", 9, "48C6D0", True, 2),
        ("Review jurisdiction-scoped cases, trends, overdue work and department performance.", 13, "FFFFFF", False, 9),
        ("SAMANVAY", 9, "48C6D0", True, 2),
        ("Multilingual guidance leads to real forms; people review and submit themselves.", 13, "FFFFFF", False, 0),
    ]
    for index, (text, size, color, bold, after) in enumerate(overview_lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.space_after = Pt(after)
        for run in paragraph.runs:
            style_run(run, size, color, bold)
    old_picture = overview.shapes[3]
    old_picture._element.getparent().remove(old_picture._element)
    if "citizen" in screenshots:
        overview_picture = overview.shapes.add_picture(str(screenshots["citizen"]), Inches(0.55), Inches(1.58), Inches(4.35), Inches(2.45))
        overview_picture.line.color.rgb = rgb("314147")
        overview_picture.line.width = Pt(1)
    add_text(overview, "SYNTHETIC PILOT  |  NO OFFICIAL DEPARTMENT SYSTEMS CONNECTED", 0.62, 4.30, 4.35, 0.42, 8.2, "9EDFE4", True, PP_ALIGN.CENTER)

    slide = body_slide(presentation, 3, "Brief about the idea")
    add_text(slide, "SamanvAI joins citizen service discovery and government case review in one role-aware workflow.", 0.43, 1.48, 9.10, 0.42, 13, MUTED)
    add_card(slide, 0.45, 2.05, 2.85, 2.55, "Fragmented journey", "Citizens need to find the right service, understand next steps and follow an application across stages.", CYAN, 11, PALE)
    add_card(slide, 3.56, 2.05, 2.85, 2.55, "Limited operational view", "Officials need comparable service, aging and SLA signals scoped to their verified state, district, taluk or department role.", "4E7DB1", 11, WHITE)
    add_card(slide, 6.67, 2.05, 2.85, 2.55, "One accountable pilot", "A citizen workspace, governance review and constrained Samanvay guide share the same case store and authorization boundaries.", CYAN, 11, PALE_GREEN)
    add_note(slide, "Pilot scope: synthetic Karnataka records and illustrative service rules; no external government portals are connected.")

    slide = body_slide(presentation, 4, "Opportunity and differentiators")
    add_card(slide, 0.45, 1.70, 2.85, 2.85, "Service continuity", "Service catalog -> application intake -> status tracking -> explicit grievance or escalation action.", CYAN, 11, PALE)
    add_card(slide, 3.56, 1.70, 2.85, 2.85, "Role-aware operations", "Server-enforced scopes keep Chief Secretary, DC, Tahsildar and department views within assigned boundaries.", "4E7DB1", 11, WHITE)
    add_card(slide, 6.67, 1.70, 2.85, 2.85, "Constrained assistance", "Six sequential workflow nodes use deterministic rules and authorized evidence. Gemini is optional phrasing, not an action authority.", CYAN, 11, PALE_GREEN)
    add_text(slide, "USP", 0.48, 4.78, 0.48, 0.27, 9, CYAN, True)
    add_text(slide, "A traceable service workflow where guidance, scope checks and human confirmation are explicit.", 1.00, 4.74, 8.50, 0.40, 12, INK, True)

    slide = body_slide(presentation, 5, "Services and capabilities")
    services = [
        ("Pension assistance", "Social Welfare | 30-day demo SLA"),
        ("Birth certificate", "Civil Registration | 14-day demo SLA"),
        ("Property registration", "Revenue | 30-day demo SLA"),
        ("Driving licence renewal", "Transport | 21-day demo SLA"),
        ("Water supply", "Public Works | 7-day demo SLA"),
    ]
    positions = [(0.45, 1.70), (3.56, 1.70), (6.67, 1.70), (2.00, 3.25), (5.10, 3.25)]
    for (title, body), (x, y) in zip(services, positions):
        add_card(slide, x, y, 2.85, 1.30, title, body, CYAN, 9.3, PALE)
    add_note(slide, "All deadlines shown are illustrative demo rules, not official processing-time commitments.", 4.98, AMBER)

    slide = body_slide(presentation, 6, "Use case: guidance to reviewed request")
    flow_nodes = [
        (0.40, "CITIZEN", "Chooses a service and shares a question"),
        (2.00, "REACT UI", "Typed forms and Samanvay guide"),
        (3.60, "FASTAPI", "Validates inputs and identity"),
        (5.20, "LANGGRAPH", "Six sequential evidence steps"),
        (6.80, "SQLALCHEMY", "Reads or writes the case store"),
        (8.40, "REVIEW", "Response, status and audit view"),
    ]
    for x, title, body in flow_nodes:
        add_flow_node(slide, x, title, body, 1.34)
    for x in (1.78, 3.38, 4.98, 6.58, 8.18):
        add_chevron(slide, x, 2.47, 0.16, 0.35)
    add_text(slide, "CHAT PATH: Citizen Service -> Policy Reasoning -> Case Resolution -> Escalation -> Sentiment/Grievance -> Governance Intelligence", 0.48, 3.60, 9.00, 0.48, 10.0, INK, True, PP_ALIGN.CENTER)
    add_card(slide, 0.58, 4.02, 4.15, 0.95, "Mutation guard", "Escalations require explicit consent or an explicit user action.", CYAN, 8.5, PALE)
    add_card(slide, 5.25, 4.02, 4.15, 0.95, "No silent submission", "Assistant navigation opens a form; the citizen reviews and submits.", "4E7DB1", 8.5, WHITE)

    slide = body_slide(presentation, 7, "Prototype screens: citizen and official workspaces")
    add_text(slide, "Screens captured from the running local pilot. Values are synthetic and may change as the demo database changes.", 0.43, 1.45, 9.10, 0.34, 10.5, MUTED)
    panels = [
        ("citizen", "Citizen workspace", "Applications first; service request options below."),
        ("map", "Regional service pulse", "Department dots over approximate pilot regions."),
        ("assistant", "Samanvay guidance", "Guided movement opens a real reviewable form."),
    ]
    for index, (key, title, caption) in enumerate(panels):
        x = 0.45 + index * 3.12
        if key in screenshots:
            slide.shapes.add_picture(str(screenshots[key]), Inches(x), Inches(1.92), Inches(2.85), Inches(1.60))
        add_text(slide, title, x, 3.68, 2.86, 0.30, 10.5, INK, True)
        add_text(slide, caption, x, 4.02, 2.86, 0.60, 8.7, MUTED)
    add_note(slide, "Samanvay currently supports English, Hindi and Kannada text guidance; spoken voice availability depends on browser/device or configured Google Cloud TTS.", 4.92)

    slide = body_slide(presentation, 8, "Architecture: current pilot and deployment target")
    add_pill(slide, 0.45, 1.42, 2.10, "IMPLEMENTED PILOT", PALE_GREEN, CYAN, 9.0)
    current = [
        (0.45, "React + TypeScript", "Vite SPA; citizen and official UI"),
        (2.30, "FastAPI", "Routes, validation and identity boundary"),
        (4.15, "LangGraph", "Six sequential workflow nodes"),
        (6.00, "SQLAlchemy", "Repository and transaction logic"),
        (7.85, "SQLite", "Synthetic local demo database"),
    ]
    for x, title, body in current:
        add_arch_node(slide, x, 1.95, 1.63, title, body, PALE)
    add_text(slide, "Optional integrations: Firebase Auth (verified tokens), Gemini/Vertex text synthesis, PostgreSQL-compatible DATABASE_URL.", 0.48, 2.86, 9.05, 0.33, 9.4, MUTED)
    add_pill(slide, 0.45, 3.48, 2.53, "CONFIGURED DEPLOYMENT PATH", PALE_AMBER, AMBER, 8.8)
    target = [
        (0.45, "GitHub Actions", "Tests and deploy"),
        (2.30, "OIDC / WIF", "Keyless Google identity"),
        (4.15, "Docker build", "Enables API and builds image"),
        (6.00, "Artifact Registry", "Commit-tagged container"),
        (7.85, "Cloud Run", "One combined app service"),
    ]
    for x, title, body in target:
        add_arch_node(slide, x, 4.00, 1.63, title, body, PALE_AMBER)
    add_text(slide, "Cloud Run deployment requires GCP setup; no live cloud deployment is claimed. Microservice/microfrontend boxes are a future migration target.", 0.48, 4.92, 9.05, 0.38, 8.7, AMBER)

    slide = body_slide(presentation, 9, "Technology stack")
    add_card(slide, 0.45, 1.62, 2.85, 2.85, "Experience", "React 19 + TypeScript\nVite build\nRecharts + Lucide\nResponsive citizen and governance views", CYAN, 10.1, PALE)
    add_card(slide, 3.56, 1.62, 2.85, 2.85, "Services and workflow", "FastAPI + Pydantic\nLangGraph: six sequential nodes\nDeterministic policy/SLA checks\nOptional Gemini/Vertex synthesis", "4E7DB1", 10.1, WHITE)
    add_card(slide, 6.67, 1.62, 2.85, 2.85, "Data and deployment", "SQLAlchemy + SQLite in local demo\nPostgreSQL-compatible connection\nFirebase token verification optional\nDocker + GitHub Actions + Cloud Run target", CYAN, 10.1, PALE_GREEN)
    add_note(slide, "No Firestore, BigQuery, Vertex RAG, external department API or notification provider is currently connected.", 4.90)

    slide = body_slide(presentation, 10, "Cost status and Google Cloud cost drivers")
    add_text(slide, "No Google Cloud project is provisioned, so there is no measured bill or defensible monthly estimate yet.", 0.45, 1.52, 9.05, 0.42, 12.0, INK, True)
    add_card(slide, 0.45, 2.10, 2.85, 2.10, "Cloud Run", "Request volume, CPU/memory, concurrency and minimum instances determine compute spend. Scale-to-zero is possible.", CYAN, 10.0, PALE)
    add_card(slide, 3.56, 2.10, 2.85, 2.10, "Cloud SQL", "Instance tier, storage, backups and connectivity are the main persistent baseline cost drivers.", "4E7DB1", 10.0, WHITE)
    add_card(slide, 6.67, 2.10, 2.85, 2.10, "Optional services", "Artifact Registry, logs, network egress, Vertex/Gemini and Cloud TTS depend on use and retention.", CYAN, 10.0, PALE_GREEN)
    add_note(slide, "Before deployment, price the chosen region, Cloud SQL tier, request profile and model usage in Google Cloud Pricing Calculator. No cost values are invented here.", 4.70, AMBER)

    slide = body_slide(presentation, 11, "Prototype snapshots")
    snapshots = [
        ("overview", "Scoped governance KPIs"),
        ("map", "Department and region counts"),
        ("form", "Citizen intake review form"),
    ]
    for index, (key, caption) in enumerate(snapshots):
        x = 0.45 + index * 3.12
        if key in screenshots:
            slide.shapes.add_picture(str(screenshots[key]), Inches(x), Inches(1.55), Inches(2.85), Inches(1.60))
        add_text(slide, caption, x, 3.28, 2.86, 0.32, 10.2, INK, True)
    add_text(slide, "Current prototype: role selector is explicitly synthetic. In production, Firebase custom claims and trusted assignments must define official scopes.", 0.48, 4.15, 9.05, 0.75, 11, MUTED)

    slide = body_slide(presentation, 12, "Prototype validation: measured checks")
    add_card(slide, 0.45, 1.65, 2.85, 1.52, "32 backend tests", "Domain, jurisdiction, retrieval, memory, API and mocked speech paths pass.", CYAN, 9.8, PALE)
    add_card(slide, 3.56, 1.65, 2.85, 1.52, "Production build", "TypeScript check and Vite production build pass.", "4E7DB1", 9.8, WHITE)
    add_card(slide, 6.67, 1.65, 2.85, 1.52, "Scope verification", "State, district, taluk and department filters are asserted at the API layer.", CYAN, 9.8, PALE_GREEN)
    add_card(slide, 0.45, 3.45, 2.85, 1.25, "Refresh behavior", "Overview refreshes on a 15-second client interval.", "4E7DB1", 9.5, WHITE)
    add_card(slide, 3.56, 3.45, 2.85, 1.25, "Map pilot", "Six named districts; approximate centroids; aggregates reconcile to scoped cases.", CYAN, 9.5, PALE)
    add_card(slide, 6.67, 3.45, 2.85, 1.25, "Not benchmarked", "No load test, p95 latency, production SLO or cloud scale test is claimed.", AMBER, 9.5, PALE_AMBER)

    slide = body_slide(presentation, 13, "Roadmap: from pilot to trusted deployment")
    roadmap = [
        (0.45, "01 / TRUST", "Provision Firebase claims; validate official roles; approve canonical policy and geography sources."),
        (2.76, "02 / DATA", "Deploy Cloud SQL with migrations, backups, private connectivity and retention controls."),
        (5.07, "03 / OPERATE", "Add request tracing, load tests, dashboards, alerts and incident/runbook ownership."),
        (7.38, "04 / EXTRACT", "Split service APIs and microfrontends only after data ownership and API contracts are defined."),
    ]
    for x, title, body in roadmap:
        add_card(slide, x, 1.78, 2.17, 2.30, title, body, CYAN if x < 5 else "4E7DB1", 9.3, PALE if x < 5 else WHITE)
    add_note(slide, "The attached Cloud Run diagram is a target architecture; current source remains one React SPA and one FastAPI runtime.", 4.68, AMBER)

    slide = body_slide(presentation, 14, "Repository, demo and product access")
    add_card(slide, 0.45, 1.72, 2.85, 2.35, "Repository", "Public GitHub repository. Repository URL is not included in this submission deck.", CYAN, 10.4, PALE)
    add_card(slide, 3.56, 1.72, 2.85, 2.35, "Demo video", "media/demo-video/samanvai-overview.mp4\n139 seconds\nEnglish voice-over", "4E7DB1", 10.4, WHITE)
    add_card(slide, 6.67, 1.72, 2.85, 2.35, "Product", "Local demo: http://127.0.0.1:8000\nNo public Cloud Run URL or custom DNS is provisioned.", CYAN, 10.4, PALE_GREEN)
    add_note(slide, "Architecture reference: architecture/cloud-run-target.html. The service split shown there is proposed, not deployed.", 4.65)

    slide = body_slide(presentation, 15, "Trust boundaries and demo limitations")
    add_bullet_stack(slide, 0.55, 1.58, 8.95, [
        ("Synthetic records", "Seeded citizen, request, feedback and SLA data are illustrative; do not enter real personal information."),
        ("No department integration", "Requests stay in the local pilot store; no external portal, email, SMS or officer notification is connected."),
        ("Human-confirmed actions", "Samanvay explains and navigates. Citizens review forms; escalations require explicit action or consent."),
        ("Verified production roles", "Firebase claims must be assigned through a trusted admin process; the demo role selector is not authentication."),
        ("Illustrative geography and policy", "Approximate pilot coordinates and demo rules are not authoritative boundaries, policies or processing-time commitments."),
    ], 10.1, 0.68, AMBER)

    presentation.save(OUT_FILE)
    print(f"Saved {OUT_FILE}")
    print(f"Slides: {len(presentation.slides)}")
    print(f"Size: {OUT_FILE.stat().st_size:,} bytes")


if __name__ == "__main__":
    build()